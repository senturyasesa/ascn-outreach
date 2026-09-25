#!/usr/bin/env python3
"""lead_engine.py — лид-движок: чаты ниши -> кто в них пишет -> проверка профиля ->
фильтры -> ИИ-оценка -> выгрузка целевых в базу рассылки.

Воркер задаётся в data/engine.json ("worker"). Может быть и рассылочным аккаунтом:
движок работает с копией его сессии и не запускается, если аккаунт заморожен за флуд. Одна задача за раз (lock). Прогресс — data/engine/job.json,
результаты — data/engine/<ниша>_chats.json и <ниша>_leads.json.

Команды (вкладка сайта и агент вызывают то же самое):
  python3 lead_engine.py worker   <сессия>                  — назначить воркер
  python3 lead_engine.py niche    <ниша> --keywords "a,b" --icp "кого ищем" [--target 150]
  python3 lead_engine.py discover <ниша>                    — найти и оценить чаты
  python3 lead_engine.py collect  <ниша> [--chats @a,@b]    — собрать и оценить лидов
  python3 lead_engine.py export   <ниша> <база> [--min 60]  — целевых -> leads.xlsx + bases.json
  python3 lead_engine.py status                             — текущая задача
"""

import os
import re
import sys
import csv
import json
import time
import argparse
import unicodedata
import urllib.request
from datetime import datetime, timezone, timedelta

os.chdir(os.path.dirname(os.path.abspath(__file__)))

import shutil
import tg
import ai_sales
import campaign as C
from telethon.sync import TelegramClient
from telethon import errors, functions
from telethon.tl.types import (User, UserStatusOnline, UserStatusOffline, UserStatusRecently,
                               UserStatusLastWeek, UserStatusLastMonth)
from telethon.tl.functions.users import GetFullUserRequest

CFG = "data/engine.json"
DIR = "data/engine"
JOB = f"{DIR}/job.json"
LOCK = f"{DIR}/lock"
PROFILES = f"{DIR}/profiles.json"      # кэш проверенных профилей: id -> данные
SOURCES = "data/lead_sources.json"      # ник -> откуда пришёл (для аналитики по источникам)
LEADS_XLSX = "data/leads.xlsx"
BASES = "data/bases.json"

# ── темп и лимиты (антифлуд) ─────────────────────────────────────────────
SLEEP_SEARCH = 3        # между поисковыми запросами
SLEEP_SAMPLE = 2        # между оценками чатов
SLEEP_CHAT = 5          # между чатами при сборе
SLEEP_ENRICH = 2.5      # между проверками профилей (тяжёлый запрос)
MAX_WAIT = 120          # FloodWait дольше — останавливаемся и сохраняем, что есть
MAX_KEYWORDS = 20
MAX_SAMPLE_CHATS = 40
SAMPLE_MSGS = 100       # выборка сообщений для оценки чата
SCAN_MSGS = 3000        # сколько последних сообщений читать при сборе
DAYS = 60               # берём только тех, кто писал за это время
MIN_CHAT_SCORE = 45     # автоотбор чатов для сбора, если не выбраны вручную
MAX_COLLECT_CHATS = 8
MAX_RESOLVE = 60        # запасной резолв по нику (строгий лимит Telegram)
PROFILE_TTL_DAYS = 30
MIN_EXPORT_SCORE = 60
AI_BATCH = 10
AI_MODEL = "openai/gpt-4o-mini"

SPAM_RE = re.compile("|".join([
    r"л[её]гк", r"от\s*18", r"18\s*\+", r"доход\s*от", r"заработ", r"подработ",
    r"став(ь|ьте)\s*\+", r"жду\s*в\s*л[си]", r"в\s*личк", r"на\s*дому",
    r"\d{3,5}\s*(в|/)\s*день", r"\d{3,5}\s*руб.*день", r"нужны\s*люди",
    r"набира[ею]\s*команд", r"свободн\w*\s*график", r"пассивн\w*\s*доход",
    r"крипт", r"трейд", r"казино", r"1\s*win", r"1\s*вин", r"промокод",
    r"бонус\s*за\s*регистр", r"пишите\s*в\s*лс", r"пиши\s*слово", r"хочешь\s*зараб",
]), re.I)

# явный продавец услуг нише по имени (в имени прописана профессия исполнителя)
SELLER_RE = re.compile("|".join([
    r"таргет", r"\bsmm\b", r"смм", r"дизайн", r"копирайт", r"ассистент", r"монтаж",
    r"маркетолог", r"сторис", r"reels", r"рилс", r"техспец", r"тех\.?\s*спец",
    r"технич\w*\s*спец", r"воронк", r"чат.?бот", r"getcourse", r"геткурс", r"тильд",
    r"taplink", r"таплинк", r"упаковк", r"визуал", r"фотограф", r"видеограф",
    r"создани\w* сайт", r"лендинг", r"\bseo\b", r"авитолог", r"ищу работу", r"резюме",
]), re.I)

DEFAULT_ICP = ("Владелец или руководитель бизнеса в этой нише, который сам принимает решения "
               "и может купить автоматизацию (ИИ-агенты для ответов клиентам, заявок, рутины).")


class EngineError(Exception):
    pass


class Flood(Exception):
    pass


# ── файлы и состояние ────────────────────────────────────────────────────
def _now():
    return datetime.now(timezone.utc)


def _load(p, d):
    try:
        return json.load(open(p, encoding="utf-8"))
    except Exception:
        return d


def _save(p, obj):
    tmp = p + ".tmp"
    json.dump(obj, open(tmp, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    os.replace(tmp, p)   # атомарно: читатель не увидит полузаписанный файл


def _slug(name):
    return re.sub(r"[^\w\-]+", "_", name.strip(), flags=re.U).strip("_") or "niche"


def chats_file(niche):
    return f"{DIR}/{_slug(niche)}_chats.json"


def leads_file(niche):
    return f"{DIR}/{_slug(niche)}_leads.json"


def cfg():
    c = _load(CFG, {})
    c.setdefault("worker", "")
    c.setdefault("niches", {})
    return c


def niche_cfg(niche):
    n = cfg()["niches"].get(niche)
    if not n:
        raise EngineError(f"ниша «{niche}» не заведена")
    return n


def job_set(**kw):
    j = _load(JOB, {})
    j.update(kw)
    j["updated"] = _now().strftime("%Y-%m-%d %H:%M:%S")
    _save(JOB, j)


def job_start(kind, niche):
    _save(JOB, {"kind": kind, "niche": niche, "stage": "старт", "progress": 0, "total": 0,
                "started": _now().strftime("%Y-%m-%d %H:%M:%S"), "done": False,
                "error": "", "summary": {}, "pid": os.getpid()})


def acquire():
    os.makedirs(DIR, exist_ok=True)
    if os.path.exists(LOCK):
        try:
            os.kill(int(open(LOCK).read().strip()), 0)
            return False                      # живой процесс держит воркер
        except Exception:
            pass                              # мёртвый lock — перехватываем
    open(LOCK, "w").write(str(os.getpid()))
    return True


def busy():
    """Идёт ли сейчас задача (живой процесс держит lock)."""
    try:
        os.kill(int(open(LOCK).read().strip()), 0)
        return True
    except Exception:
        return False


def release():
    try:
        os.remove(LOCK)
    except Exception:
        pass


# ── мелкие проверки ──────────────────────────────────────────────────────
def full_name(u):
    return " ".join(filter(None, [getattr(u, "first_name", None), getattr(u, "last_name", None)]))


def junk_name(name):
    vis = [ch for ch in str(name or "") if unicodedata.category(ch)[0] in "LN"]
    return len(vis) < 2


def is_spam(*texts):
    return any(t and SPAM_RE.search(t) for t in texts)


def status_label(st):
    if isinstance(st, (UserStatusOnline, UserStatusRecently)):
        return "недавно"
    if isinstance(st, UserStatusLastWeek):
        return "на неделе"
    if isinstance(st, UserStatusLastMonth):
        return "в этом месяце"
    if isinstance(st, UserStatusOffline) and st.was_online:
        days = (_now() - st.was_online).days
        return "недавно" if days <= 7 else ("в этом месяце" if days <= 45 else "давно")
    return "давно"      # UserStatusEmpty / нет данных — «был давно»


def _key(n):
    return str(n or "").lstrip("@").strip().lower()


def known_nicks():
    """Все ники, которым уже писали или которые уже лежат в базе/блэклисте."""
    known = set()
    if os.path.exists("data/sent_log.csv"):
        for r in csv.DictReader(open("data/sent_log.csv", encoding="utf-8-sig")):
            known.add(_key(r.get("ник")))
    try:
        import openpyxl
        ws = openpyxl.load_workbook(LEADS_XLSX, read_only=True).active
        hdr = [c.value for c in next(ws.iter_rows(max_row=1))]
        ni = hdr.index("ник")
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[ni]:
                known.add(_key(row[ni]))
    except Exception:
        pass
    for nicks in _load(BASES, {}).values():
        known.update(_key(n) for n in nicks)
    known.update(_key(n) for n in _load("data/blacklist.json", {}))
    known.update(_key(n) for n in _load("data/invalid.json", {}))
    known.discard("")
    return known


# ── Telegram ─────────────────────────────────────────────────────────────
def open_worker():
    """Воркер может быть и рассылочным аккаунтом. Чтобы не ловить «database is locked»
    с демоном рассылки, работаем со СВОЕЙ копией сессии (тот же ключ, другой файл)."""
    w = cfg()["worker"].strip()
    if not w:
        raise EngineError("воркер не назначен")
    src = f"data/{w}.session"
    if not os.path.exists(src):
        raise EngineError(f"нет файла сессии {src}")
    try:
        import spambot
        if spambot.is_frozen(w):
            raise EngineError(f"«{w}» заморожен за флуд — парсинг пропущен, чтобы не добить аккаунт")
    except ImportError:
        pass
    os.makedirs("data/_engine", exist_ok=True)
    dst = f"data/_engine/{w}"
    for suf in (".session-journal", ".session-wal", ".session-shm"):
        if os.path.exists(dst + suf):
            os.remove(dst + suf)
    shutil.copy2(src, dst + ".session")
    api_id, api_hash = tg.credentials()
    c = TelegramClient(dst, api_id, api_hash, proxy=tg.proxy_for(w),
                       flood_sleep_threshold=60, **C.device_for(w))
    c.connect()
    if not c.is_user_authorized():
        c.disconnect()
        raise EngineError(f"воркер «{w}» не авторизован")
    return c


def call(fn, *a, **k):
    try:
        return fn(*a, **k)
    except errors.FloodWaitError as e:
        if e.seconds > MAX_WAIT:
            raise Flood(e.seconds)
        time.sleep(e.seconds + 2)
        return fn(*a, **k)


# ── 1. поиск и оценка чатов ──────────────────────────────────────────────
def score_chat(c, ch, kw):
    row = {"username": ch.username, "title": ch.title or "", "id": ch.id, "keyword": kw,
           "members": getattr(ch, "participants_count", None) or 0}
    try:
        msgs = call(c.get_messages, ch, limit=SAMPLE_MSGS)
    except Flood:
        raise
    except Exception as e:
        row.update(score=0, note=f"не читается ({type(e).__name__})")
        return row
    users, spam, texts, last = {}, 0, 0, None
    for m in msgs:
        if m.date and (last is None or m.date > last):
            last = m.date
        u = m.sender
        if not isinstance(u, User) or u.bot:
            continue
        t = m.message or ""
        if t:
            texts += 1
            if is_spam(t):
                spam += 1
        users[u.id] = u
    idle = (_now() - last).days if last else 999
    posters = len(users)
    sellers = sum(1 for u in users.values() if SELLER_RE.search(full_name(u)))
    spam_r = spam / max(texts, 1)
    sell_r = sellers / max(posters, 1)
    score = (30 if idle <= 3 else 20 if idle <= 14 else 5 if idle <= 45 else 0)
    score += min(30, posters)
    score += 20 * (1 - spam_r) + 20 * (1 - sell_r)
    note = ""
    if idle > DAYS:        # сбор берёт только свежих авторов — из мёртвого чата будет ноль
        score, note = min(score, 20), f"мёртвый: {idle} дн без сообщений"
    row.update(score=round(score), idle_days=idle, posters=posters,
               spam_pct=round(100 * spam_r), sellers_pct=round(100 * sell_r), note=note)
    return row


def discover(niche, max_chats=MAX_SAMPLE_CHATS):
    n = niche_cfg(niche)
    kws = [k.strip() for k in n.get("keywords", []) if k.strip()][:MAX_KEYWORDS]
    if not kws:
        raise EngineError("у ниши нет ключевых слов")
    c = open_worker()
    found = {}
    try:
        for i, kw in enumerate(kws, 1):
            job_set(stage=f"поиск чатов: «{kw}»", progress=i, total=len(kws))
            res = call(c, functions.contacts.SearchRequest(q=kw, limit=50))
            for ch in res.chats:
                # только группы с обсуждением: из канала пишущих не достать
                if getattr(ch, "megagroup", False) and getattr(ch, "username", None):
                    found.setdefault(ch.id, (ch, kw))
            time.sleep(SLEEP_SEARCH)

        items = list(found.values())[:max_chats]
        prev = {r["username"]: r for r in _load(chats_file(niche), [])}
        rows = []
        for i, (ch, kw) in enumerate(items, 1):
            job_set(stage=f"оценка чата: {ch.title}", progress=i, total=len(items))
            r = score_chat(c, ch, kw)
            r["selected"] = (False if r.get("idle_days", 0) > DAYS
                             else prev.get(ch.username, {}).get("selected", r["score"] >= MIN_CHAT_SCORE))
            rows.append(r)
            time.sleep(SLEEP_SAMPLE)
    finally:
        c.disconnect()
    rows.sort(key=lambda r: -r["score"])
    _save(chats_file(niche), rows)
    return {"найдено групп": len(found), "оценено": len(rows),
            "рекомендовано": sum(1 for r in rows if r["selected"])}


# ── 2. сбор, проверка профиля, фильтры, ИИ-оценка ───────────────────────
def _enrich(c, u, username, resolves):
    """Проверка профиля одним GetFullUser. u — объект из сообщений (без резолва ника)."""
    try:
        full = call(c, GetFullUserRequest(u))
    except Flood:
        raise
    except Exception:
        if resolves[0] >= MAX_RESOLVE:
            return None
        resolves[0] += 1
        try:
            full = call(c, GetFullUserRequest(username))
        except Flood:
            raise
        except Exception:
            return None
    fu = full.full_user
    uu = full.users[0] if full.users else u
    return {
        "bio": (fu.about or "")[:400],
        "channel": bool(getattr(fu, "personal_channel_id", None)),
        "business": bool(getattr(fu, "business_intro", None) or getattr(fu, "business_location", None)
                         or getattr(fu, "business_work_hours", None)),
        "paid": int(getattr(fu, "send_paid_messages_stars", 0) or getattr(uu, "send_paid_messages_stars", 0) or 0),
        "premium_only": bool(getattr(uu, "contact_require_premium", False)),
        "premium": bool(getattr(uu, "premium", False)),
        "seen": status_label(getattr(uu, "status", None)),
        "checked": _now().strftime("%Y-%m-%d"),
    }


def drop_reason(d):
    if junk_name(d.get("name")):
        return "пустое/невидимое имя"
    if d.get("paid"):
        return "платные сообщения"
    if d.get("premium_only"):
        return "пишут только премиум"
    if d.get("seen") == "давно":
        return "давно не заходил"
    if SELLER_RE.search(d.get("name") or ""):
        return "продавец услуг (по имени)"
    return ""


def ai_score(batch, icp, niche):
    """ИИ-оценка пачки кандидатов. -> {ник: (класс, балл, причина)}."""
    key = ai_sales._cfg().get("key")
    if not key:
        return {}
    sysmsg = (
        "Ты оцениваешь лидов для холодной Telegram-рассылки ASCN (платформа ИИ-агентов: "
        "автоответы клиентам, обработка заявок, рутина бизнеса). "
        f"Ниша: «{niche}». Кого ищем: {icp}\n"
        "Для каждого кандидата выбери класс:\n"
        "лпр — владелец, руководитель, эксперт или автор своего дела, который сам может купить "
        "автоматизацию для СВОЕГО бизнеса;\n"
        "продавец — сам продаёт услуги этой нише (техспец, таргетолог, SMM, чат-боты, дизайн, "
        "ассистент, менеджер в найме, фрилансер) — нам не подходит;\n"
        "мусор — бот, спам, случайный человек, не из ниши.\n"
        "ВАЖНО: кто сам делает ИИ, нейросети, ботов, автоматизацию или внедряет их другим — "
        "это продавец (наш конкурент), даже если он эксперт или основатель.\n"
        "score 0-100 — насколько вероятно, что это наш покупатель. Если есть только имя "
        "(нет bio и содержательных сообщений) — score не выше 40.\n"
        "reason — по-русски, до 12 слов, на чём основано решение.\n"
        'Верни ТОЛЬКО JSON-массив: [{"u":"ник","class":"лпр|продавец|мусор","score":0,"reason":"..."}]')
    items = [{"u": d["username"], "имя": d.get("name", ""), "bio": d.get("bio", ""),
              "свой_канал": d.get("channel", False), "бизнес_аккаунт": d.get("business", False),
              "чат": d.get("chat", ""), "сообщения": d.get("msgs", [])[:3]} for d in batch]
    try:
        body = json.dumps({"model": AI_MODEL, "temperature": 0,
                           "messages": [{"role": "system", "content": sysmsg},
                                        {"role": "user", "content": json.dumps(items, ensure_ascii=False)}]
                           }).encode("utf-8")
        req = urllib.request.Request(ai_sales.API_URL, data=body, headers={
            "Authorization": "Bearer " + key, "Content-Type": "application/json"})
        with urllib.request.urlopen(req, timeout=60) as resp:
            out = json.load(resp)["choices"][0]["message"]["content"]
        out = re.sub(r"^```(?:json)?|```$", "", out.strip(), flags=re.M).strip()
        data = json.loads(out)
        if isinstance(data, dict):     # {"items":[...]} вместо массива
            data = next((v for v in data.values() if isinstance(v, list)), [])
        res = {}
        for r in data:
            cls = str(r.get("class", "")).lower()
            cls = cls if cls in ("лпр", "продавец", "мусор") else "мусор"
            res[_key(r.get("u"))] = (cls, max(0, min(100, int(r.get("score", 0) or 0))),
                                     str(r.get("reason", ""))[:120])
        return res
    except Exception as e:
        print("ai_score err:", type(e).__name__, e)
        return {}


SKIPPED = "не проверен (лимит на прогон)"


def _keepable(rows):
    """В файл лидов не кладём тех, до кого не дошла проверка, — подберём в следующий раз."""
    return [d for d in rows if d.get("drop") != SKIPPED]


def collect(niche, chats=None):
    n = niche_cfg(niche)
    target = int(n.get("target", 150) or 150)
    icp = (n.get("icp") or "").strip() or DEFAULT_ICP
    rows = _load(chats_file(niche), [])
    if chats:
        sel = ["@" + _key(x) for x in chats if _key(x)]
    else:
        sel = ["@" + r["username"] for r in rows if r.get("selected")][:MAX_COLLECT_CHATS]
    if not sel:
        raise EngineError("нет выбранных чатов: сначала discover или передай --chats")

    known = known_nicks()
    prev = {_key(d["username"]): d for d in _load(leads_file(niche), [])}
    known.update(prev)                                   # не собираем повторно
    per_chat_cap = max(20, int(target * 2.5 / len(sel)))
    cutoff = _now() - timedelta(days=DAYS)
    cand, objs, spammers = {}, {}, set()
    stats = {"чатов": len(sel), "прочитано сообщений": 0}

    c = open_worker()
    try:
        for i, ch in enumerate(sel, 1):
            job_set(stage=f"читаю чат {ch}", progress=i, total=len(sel))
            got = 0
            try:
                ent = call(c.get_entity, ch)
                for m in c.iter_messages(ent, limit=SCAN_MSGS):
                    stats["прочитано сообщений"] += 1
                    if m.date and m.date < cutoff:
                        break
                    u = m.sender
                    if (not isinstance(u, User) or u.bot or u.deleted or u.is_self
                            or not u.username):
                        continue
                    t = m.message or ""
                    if u.id in spammers:
                        continue
                    if is_spam(t, full_name(u)):
                        spammers.add(u.id)
                        cand.pop(u.id, None)
                        continue
                    if _key(u.username) in known:
                        continue
                    d = cand.get(u.id)
                    if d is None:
                        if got >= per_chat_cap:
                            continue
                        d = cand[u.id] = {"id": u.id, "username": u.username, "name": full_name(u),
                                          "chat": ch, "msgs": [], "n": 0,
                                          "last": m.date.strftime("%Y-%m-%d") if m.date else "",
                                          "paid": int(getattr(u, "send_paid_messages_stars", 0) or 0),
                                          "premium_only": bool(getattr(u, "contact_require_premium", False))}
                        objs[u.id] = u
                        got += 1
                    d["n"] += 1
                    if t and len(t) > 15 and len(d["msgs"]) < 3:
                        d["msgs"].append(t[:300])
            except errors.FloodWaitError as e:
                raise Flood(e.seconds)
            except Flood:
                raise
            except Exception as e:
                print(f"чат {ch}: {type(e).__name__} {e}")
            time.sleep(SLEEP_CHAT)

        # до проверки профиля отсекаем очевидное — не тратим тяжёлые запросы
        for d in cand.values():
            d["drop"] = drop_reason(d)

        alive = sorted((d for d in cand.values() if not d["drop"]), key=lambda d: -d["n"])
        to_check = alive[:max(20, target * 2)]
        profiles = _load(PROFILES, {})
        fresh = (_now() - timedelta(days=PROFILE_TTL_DAYS)).strftime("%Y-%m-%d")
        resolves = [0]
        for i, d in enumerate(to_check, 1):
            job_set(stage="проверяю профили", progress=i, total=len(to_check))
            p = profiles.get(str(d["id"]))
            if not p or p.get("checked", "") < fresh:
                p = _enrich(c, objs[d["id"]], d["username"], resolves)
                time.sleep(SLEEP_ENRICH)
                if p:
                    profiles[str(d["id"])] = p
            if p:
                d.update(p)
                d["drop"] = drop_reason(d)
            else:
                d["drop"] = "профиль не открылся"
        for d in alive[len(to_check):]:
            d["drop"] = SKIPPED
        _save(PROFILES, profiles)
    except Flood as e:
        _save(leads_file(niche), list(prev.values()) + _keepable(cand.values()))
        raise EngineError(f"Telegram притормозил воркер на {e.args[0]} сек. "
                          f"Собранное сохранено, запусти сбор позже.")
    finally:
        c.disconnect()

    # ИИ-оценка — только тех, кто прошёл фильтры
    scored = [d for d in cand.values() if not d["drop"]]
    for i in range(0, len(scored), AI_BATCH):
        batch = scored[i:i + AI_BATCH]
        job_set(stage="ИИ-оценка", progress=min(i + AI_BATCH, len(scored)), total=len(scored))
        res = ai_score(batch, icp, niche)
        for d in batch:
            cls, sc, why = res.get(_key(d["username"]), ("?", 0, "ИИ не ответил — глянь вручную"))
            d.update({"class": cls, "score": sc, "reason": why})

    allrows = list(prev.values()) + _keepable(cand.values())
    _save(leads_file(niche), allrows)

    drops = {}
    for d in cand.values():
        if d["drop"]:
            drops[d["drop"]] = drops.get(d["drop"], 0) + 1
    classes = {}
    for d in scored:
        classes[d.get("class", "?")] = classes.get(d.get("class", "?"), 0) + 1
    good = sum(1 for d in scored if d.get("class") == "лпр" and d.get("score", 0) >= MIN_EXPORT_SCORE)
    stats.update({"новых кандидатов": len(cand), "отсеяно": drops, "ИИ-оценка": classes,
                  f"целевых (лпр, балл ≥{MIN_EXPORT_SCORE})": good})
    return stats


# ── 3. выгрузка в рассылку ───────────────────────────────────────────────
def export(niche, base, min_score=MIN_EXPORT_SCORE):
    import openpyxl
    rows = _load(leads_file(niche), [])
    known = known_nicks()
    picks = [d for d in rows if not d.get("drop") and d.get("class") == "лпр"
             and d.get("score", 0) >= min_score and not d.get("exported")
             and _key(d["username"]) not in known]
    picks.sort(key=lambda d: -d.get("score", 0))
    if not picks:
        return {"добавлено": 0}

    wb = openpyxl.load_workbook(LEADS_XLSX)
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(hdr)}
    last = 0
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, col["№"]).value if "№" in col else None
        if isinstance(v, (int, float)):
            last = max(last, int(v))
    for d in picks:
        r = ws.max_row + 1
        last += 1
        vals = {"№": last, "тир": "A" if d["score"] >= 80 else "B",
                "ник": "@" + d["username"], "имя": d.get("name", ""),
                "суть клиента": d.get("reason", ""),
                "его сообщения из чата": "\n".join(d.get("msgs", []))}
        for h, v in vals.items():
            if h in col:
                ws.cell(r, col[h], v)
    tmp = LEADS_XLSX + ".tmp.xlsx"
    wb.save(tmp)
    os.replace(tmp, LEADS_XLSX)          # атомарно: демон не прочтёт полфайла

    bases = _load(BASES, {})
    cur = bases.setdefault(base, [])
    have = {_key(x) for x in cur}
    cur += ["@" + d["username"] for d in picks if _key(d["username"]) not in have]
    _save(BASES, bases)

    src = _load(SOURCES, {})
    today = _now().strftime("%Y-%m-%d")
    for d in picks:
        src["@" + d["username"]] = {"chat": d.get("chat", ""), "niche": niche, "base": base,
                                    "score": d.get("score"), "added": today}
        d["exported"] = today
    _save(SOURCES, src)
    _save(leads_file(niche), rows)
    return {"добавлено": len(picks), "база": base}


# ── запуск ───────────────────────────────────────────────────────────────
def run(kind, niche, fn, *a):
    if not acquire():
        print("воркер занят другой задачей")
        return 2
    job_start(kind, niche)
    try:
        summary = fn(*a)
        job_set(stage="готово", done=True, summary=summary)
        print(json.dumps(summary, ensure_ascii=False, indent=1))
        return 0
    except EngineError as e:
        job_set(stage="остановлено", done=True, error=str(e))
        print("ошибка:", e)
        return 1
    except Exception as e:
        job_set(stage="сбой", done=True, error=f"{type(e).__name__}: {e}")
        raise
    finally:
        release()


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("worker"); s.add_argument("session")
    s = sub.add_parser("niche"); s.add_argument("name"); s.add_argument("--keywords", default="")
    s.add_argument("--icp", default=""); s.add_argument("--target", type=int, default=0)
    s = sub.add_parser("discover"); s.add_argument("niche")
    s.add_argument("--max-chats", type=int, default=MAX_SAMPLE_CHATS)
    s = sub.add_parser("collect"); s.add_argument("niche"); s.add_argument("--chats", default="")
    s = sub.add_parser("export"); s.add_argument("niche"); s.add_argument("base")
    s.add_argument("--min", type=int, default=MIN_EXPORT_SCORE)
    sub.add_parser("status")
    a = ap.parse_args()
    os.makedirs(DIR, exist_ok=True)

    if a.cmd == "worker":
        c = cfg()
        c["worker"] = a.session
        _save(CFG, c); print("воркер:", a.session); return 0
    if a.cmd == "niche":
        c = cfg()
        n = c["niches"].setdefault(a.name, {"keywords": [], "icp": "", "target": 150})
        if a.keywords:
            n["keywords"] = [k.strip() for k in a.keywords.split(",") if k.strip()]
        if a.icp:
            n["icp"] = a.icp
        if a.target:
            n["target"] = a.target
        _save(CFG, c); print(json.dumps(n, ensure_ascii=False, indent=1)); return 0
    if a.cmd == "status":
        print(json.dumps(_load(JOB, {}), ensure_ascii=False, indent=1)); return 0
    if a.cmd == "discover":
        return run("discover", a.niche, discover, a.niche, a.max_chats)
    if a.cmd == "collect":
        chats = [x for x in a.chats.split(",") if x.strip()] or None
        return run("collect", a.niche, collect, a.niche, chats)
    if a.cmd == "export":
        return run("export", a.niche, export, a.niche, a.base, a.min)


if __name__ == "__main__":
    sys.exit(main() or 0)
