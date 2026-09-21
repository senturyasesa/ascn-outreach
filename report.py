#!/usr/bin/env python3
"""report.py — вечерний отчёт по дневной рассылке в Telegram-бот.
Запуск по cron в 19:00 МСК (16:00 UTC). Шлёт в ту же группу, что уведомления."""

import os
import csv
import json
import datetime
import urllib.parse
import urllib.request

os.chdir(os.path.dirname(os.path.abspath(__file__)))
import daily_sender as D

BOT = os.environ.get("TG_BOT_TOKEN", "")
CHAT = os.environ.get("TG_CHAT_ID", "")
LOG = "data/sent_log.csv"


def send(text):
    data = urllib.parse.urlencode({
        "chat_id": CHAT, "text": text, "disable_web_page_preview": "true"
    }).encode()
    urllib.request.urlopen(
        f"https://api.telegram.org/bot{BOT}/sendMessage", data=data, timeout=20)


def main():
    utc_today = datetime.datetime.utcnow().strftime("%Y-%m-%d")
    ok = fail = 0
    by_acc = {}
    done = set()
    if os.path.exists(LOG):
        for r in csv.DictReader(open(LOG, encoding="utf-8-sig")):
            st = r.get("статус")
            if st == "ok":
                done.add(r["ник"])
            if str(r.get("время", ""))[:10] == utc_today:
                if st == "ok":
                    ok += 1
                    a = r.get("аккаунт", "")
                    by_acc[a] = by_acc.get(a, 0) + 1
                elif st == "fail":
                    fail += 1

    replies = {}
    if os.path.exists("data/replies.json"):
        try:
            replies = json.load(open("data/replies.json", encoding="utf-8"))
        except Exception:
            pass

    import openpyxl
    ws = openpyxl.load_workbook("data/leads.xlsx").active
    tot = sum(1 for r in range(2, ws.max_row + 1) if ws.cell(r, 3).value)

    target, _ = D.load_plan()
    now = D.now_msk()

    m = f"📊 Отчёт по рассылке · {now.strftime('%d.%m %H:%M')} МСК\n\n"
    m += f"✅ Отправлено сегодня: {ok} из {target}\n"
    if fail:
        m += f"⚠️ Не прошло попыток: {fail} (обошли другими акками)\n"
    m += f"💬 Ответили всего: {len(replies)}\n"
    m += f"📇 База: {len(done)}/{tot} отправлено, осталось {tot - len(done)}\n"

    if by_acc:
        m += "\nПо аккаунтам сегодня:\n"
        for a, n in sorted(by_acc.items(), key=lambda x: -x[1]):
            m += f"  · {a}: {n}\n"

    if replies:
        m += "\nПоследние ответы:\n"
        for nick, v in list(replies.items())[-5:]:
            txt = (v.get("текст", "") if isinstance(v, dict) else str(v))[:60]
            m += f"  · {nick}: {txt}\n"

    send(m)
    print("отчёт отправлен в Telegram")


if __name__ == "__main__":
    main()
